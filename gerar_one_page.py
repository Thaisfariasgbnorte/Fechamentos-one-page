"""Gera a ONE PAGE de Pós-Vendas (index.html) a partir da planilha mãe.

Uso:
    python3 gerar_one_page.py INDICADORES_POS_VENDA_planilha_mae.xlsx [saida.html]

Lê os valores já calculados pelo Excel (salve a planilha no Excel antes de rodar)
e embute os dados em template.html.
"""
import json
import sys
from pathlib import Path

import openpyxl

MESES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
MESES_EXTENSO = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO",
                 "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]

# aba -> (chave, rótulo, unidade, como consolidar)
INDICADORES = [
    ("FATURAMENTO", "fat", "Faturamento", "R$", "soma"),
    ("MARGEM BRUTA", "margem", "Margem bruta", "%", "media_pond_fat"),
    ("PASSAGENS", "pass", "Passagens", "qtd", "soma"),
    ("TICKET MÉDIO", "ticket", "Ticket médio", "R$", "media_pond_pass"),
    ("NPS", "nps", "NPS", "pts", "media"),
    ("FFV", "ffv", "FFV", "%", "media"),
    ("DID", "did", "DID", "%", "media"),
]


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def bloco(ws, linha_ini, n_lojas):
    """Lê lojas (col B) e 12 meses (cols C:N) a partir de linha_ini; última linha = consolidado."""
    lojas, cons = {}, None
    for r in range(linha_ini, linha_ini + n_lojas + 1):
        nome = ws.cell(r, 2).value
        vals = [num(ws.cell(r, c).value) for c in range(3, 15)]
        if r == linha_ini + n_lojas:
            cons = vals
        else:
            lojas[nome] = vals
    return lojas, cons


def main(xlsx, saida):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    cfg = wb["CONFIGURAÇÃO"]
    lojas = [cfg.cell(r, 3).value for r in range(12, 21) if cfg.cell(r, 3).value]
    n = len(lojas)

    fontes = {}
    for r in range(30, 37):
        fontes[cfg.cell(r, 2).value] = cfg.cell(r, 3).value

    dados = {
        "empresa": cfg["C6"].value,
        "ano": cfg["C7"].value,
        "mesExport": MESES.index(str(cfg["C9"].value).strip().lower()) if cfg["C9"].value else None,
        "periodo1": cfg["C23"].value,
        "periodo2": cfg["C24"].value,
        "inicioPeriodo2": MESES.index(cfg["C25"].value) if cfg["C25"].value in MESES else None,
        "faixaVerde": cfg["C39"].value or 1,
        "faixaAmarela": cfg["C40"].value or 0.9,
        "lojas": lojas,
        "meses": MESES,
        "indicadores": [],
    }

    for aba, chave, rotulo, unidade, consol in INDICADORES:
        ws = wb[aba]
        real, real_cons = bloco(ws, 8, n)
        meta, meta_cons = bloco(ws, 21, n)
        comentarios = {ws.cell(r, 2).value: ws.cell(r, 3).value
                       for r in range(47, 47 + n + 1) if ws.cell(r, 3).value}
        dados["indicadores"].append({
            "chave": chave, "rotulo": rotulo, "unidade": unidade, "consolidacao": consol,
            "fonte": fontes.get(rotulo, ""),
            "real": real, "realCons": real_cons, "meta": meta, "metaCons": meta_cons,
            "comentarios": comentarios,
        })

    # Oficina (export Linx/NBS do mês atual)
    ro = wb["RESUMO OFICINA"]
    cols = ["passagens", "mo", "objMo", "pctMo", "pecas", "objPecas", "pctPecas", "total",
            "projecao", "ticket", "lbPecas", "margemPecas", "pecasMo", "colaboradores"]
    oficina = {}
    for r in range(7, 7 + n + 2):
        nome = ro.cell(r, 2).value
        if not nome:
            continue
        oficina[nome] = {k: num(ro.cell(r, 3 + i).value) for i, k in enumerate(cols)}
    sem_loja = [ro.cell(21, c).value for c in range(3, 20) if ro.cell(21, c).value]

    ranking = []
    for r in range(25, ro.max_row + 1):
        nome, loja = ro.cell(r, 3).value, ro.cell(r, 4).value
        total = num(ro.cell(r, 8).value)
        if not nome or not loja or loja == "SEM LOJA" or not total:
            continue
        ranking.append({
            "origem": ro.cell(r, 2).value, "nome": nome, "loja": loja,
            "passagens": num(ro.cell(r, 5).value), "mo": num(ro.cell(r, 6).value),
            "pecas": num(ro.cell(r, 7).value), "total": total,
            "ticket": num(ro.cell(r, 9).value), "margem": num(ro.cell(r, 10).value),
        })
    ranking.sort(key=lambda x: -x["total"])
    dados["oficina"] = {"porLoja": oficina, "semLoja": sem_loja, "ranking": ranking}

    # Estoque de peças (aba compra de peças: B=mês, D:G saldos, L=loja one page)
    cp = wb["compra de peças"]
    estoque = []
    for r in range(8, cp.max_row + 1):
        mes, loja = cp.cell(r, 2).value, cp.cell(r, 12).value
        if not mes or not loja or str(mes).strip().upper() not in MESES_EXTENSO:
            continue
        estoque.append({
            "mes": MESES_EXTENSO.index(str(mes).strip().upper()), "loja": loja,
            "inicial": num(cp.cell(r, 4).value), "entradas": num(cp.cell(r, 5).value),
            "saidas": num(cp.cell(r, 6).value), "final": num(cp.cell(r, 7).value),
        })
    dados["estoque"] = estoque

    # Qualidade montadora
    qm = wb["QUALIDADE MONTADORA"]
    qualidade = []
    for r in range(7, qm.max_row + 1):
        if not qm.cell(r, 3).value:
            continue
        qualidade.append({
            "marca": qm.cell(r, 2).value, "loja": qm.cell(r, 3).value,
            "programa": qm.cell(r, 4).value, "meta": num(qm.cell(r, 5).value),
            "resultado": num(qm.cell(r, 6).value), "status": qm.cell(r, 7).value,
            "obs": qm.cell(r, 8).value,
        })
    dados["qualidade"] = {"periodo": qm["C4"].value, "itens": qualidade}

    template = (Path(__file__).parent / "template.html").read_text(encoding="utf-8")
    html = template.replace("/*__DADOS__*/null", json.dumps(dados, ensure_ascii=False))
    Path(saida).write_text(html, encoding="utf-8")
    print(f"OK: {saida}  ({len(lojas)} lojas, {len(ranking)} colaboradores, {len(estoque)} linhas de estoque)")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "index.html")
